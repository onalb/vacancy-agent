"""Turn a rent-roll workbook into plan counts. Resident names are never returned."""

from __future__ import annotations

import calendar
import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

_UNIT_RE = re.compile(r"^\d+-\d+")
_ASOF_RE = re.compile(r"(?i)rent roll as of:\s*(.+)")
_MONTHS = {
    "jan": 1,
    "feb": 2,
    "mar": 3,
    "apr": 4,
    "may": 5,
    "jun": 6,
    "jul": 7,
    "aug": 8,
    "sep": 9,
    "oct": 10,
    "nov": 11,
    "dec": 12,
}


def parse_rent_roll(path: str | Path) -> dict:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if "Rent Roll" in workbook.sheetnames:
            return _rediq(workbook["Rent Roll"])
        sheet = workbook[workbook.sheetnames[0]]
        return _realpage(sheet)
    finally:
        workbook.close()


def _realpage(sheet) -> dict:
    rows = list(sheet.iter_rows(values_only=True))
    property_name = ""
    as_of = ""
    header_at = None
    for index, row in enumerate(rows[:12]):
        text = str(row[0] or "").strip()
        if not text:
            continue
        if text.lower() == "rent roll" or text.lower() == "unit details":
            continue
        parsed = _as_of(text)
        if parsed and not as_of:
            as_of = parsed
            continue
        if header_at is None and "bldg-unit" in text.lower():
            header_at = index
            continue
        if not property_name:
            property_name = text
    if header_at is None:
        raise ValueError("This workbook has no unit header.")
    labels = [str(cell or "").strip().lower() for cell in rows[header_at]]
    plan_at = _index(labels, "unit type")
    status_at = _index(labels, "unit status")
    latest: dict[str, tuple[str, int, str]] = {}
    for row in rows[header_at + 1 :]:
        unit_id = row[0]
        if not isinstance(unit_id, str) or not _UNIT_RE.match(unit_id.strip()):
            continue
        status = str(row[status_at] or "").strip().lower()
        if any(word in status for word in ("model", "excluded", "non-rev", "nonrev")):
            latest.pop(unit_id.strip(), None)
            continue
        if "vacant" in status:
            occupancy = "vacant"
        elif "occup" in status or "notice" in status:
            occupancy = "occupied"
        else:
            continue
        plan, renovated = _split_plan(str(row[plan_at] or ""))
        if plan:
            latest[unit_id.strip()] = (plan, renovated, occupancy)
    return _aggregate(property_name, as_of, list(latest.values()))


def _rediq(sheet) -> dict:
    rows = list(sheet.iter_rows(values_only=True))
    property_name = ""
    as_of = ""
    for row in rows[:8]:
        for cell in row:
            text = str(cell or "").strip()
            match = _ASOF_RE.search(text)
            if match:
                as_of = _as_of(match.group(1)) or as_of
            elif "units)" in text.lower() and not property_name:
                property_name = text.split("(")[0].strip()
    header_at = next(
        index
        for index, row in enumerate(rows[:20])
        if any(str(cell or "").strip() == "UnitID" for cell in row)
    )
    labels = [str(cell or "").strip() for cell in rows[header_at]]
    index = {name: position for position, name in enumerate(labels) if name}
    plan_at = index["PlanID"]
    status_at = index["OccStatus"]
    unit_at = index["UnitID"]
    reno_at = index.get("RenStatus")
    units: list[tuple[str, int, str]] = []
    for row in rows[header_at + 1 :]:
        unit_id = row[unit_at]
        if not isinstance(unit_id, str) or not unit_id.strip() or unit_id.strip() == "UnitID":
            continue
        status = str(row[status_at] or "").strip().lower()
        if any(word in status for word in ("non", "model", "excluded", "down")):
            continue
        if "vacant" in status:
            occupancy = "vacant"
        elif "occup" in status:
            occupancy = "occupied"
        else:
            continue
        plan, renovated = _split_plan(str(row[plan_at] or ""))
        if reno_at is not None and "reno" in str(row[reno_at] or "").lower():
            renovated = 1
        if plan:
            units.append((plan, renovated, occupancy))
    return _aggregate(property_name, as_of, units)


def _aggregate(property_name: str, as_of: str, units: list[tuple[str, int, str]]) -> dict:
    grouped: dict[tuple[str, int], list[int]] = {}
    for plan, renovated, occupancy in units:
        bucket = grouped.setdefault((plan, renovated), [0, 0])
        bucket[0] += 1
        bucket[1] += 1 if occupancy == "vacant" else 0
    plans = [
        {"plan": plan, "renovated": renovated, "units": counts[0], "vacant": counts[1]}
        for (plan, renovated), counts in sorted(grouped.items())
    ]
    return {
        "property": property_name,
        "as_of": as_of,
        "units": sum(item["units"] for item in plans),
        "vacant": sum(item["vacant"] for item in plans),
        "plans": plans,
    }


def _split_plan(raw: str) -> tuple[str, int]:
    text = " ".join(raw.split())
    lowered = text.lower()
    for suffix in (" renovated", " reno"):
        if lowered.endswith(suffix):
            return text[: -len(suffix)].strip(), 1
    return text, 0


def _as_of(text: str) -> str:
    cleaned = " ".join(text.replace(",", " ").split())
    for fmt in ("%B %d %Y", "%b %d %Y", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(cleaned, fmt).date().isoformat()
        except ValueError:
            continue
    match = re.fullmatch(r"([A-Za-z]+)\s+(\d{4})", cleaned)
    if not match:
        return ""
    month = _MONTHS.get(match.group(1)[:3].lower())
    if month is None:
        return ""
    year = int(match.group(2))
    return date(year, month, calendar.monthrange(year, month)[1]).isoformat()


def _index(labels: list[str], name: str) -> int:
    for position, label in enumerate(labels):
        if label == name:
            return position
    raise ValueError(f"Missing column {name}")
